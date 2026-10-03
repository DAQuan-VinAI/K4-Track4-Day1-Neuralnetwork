"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
import re
import shutil
import statistics
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
           "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
           "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
           "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
MAX_ROWS = 60  # mẫu có sẵn công thức cho dòng 2..61; Seeds/Summary cũng chỉ tham chiếu vùng này

# Cách ghi trong bảng (theo danh sách chọn của mẫu)
OPTIMIZER_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{result['cfg']['exp_id']}.json"
    payload = {k: result[k] for k in ("cfg", "history", "summary")}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results = []
    for path in Path(results_dir).glob("*.json"):
        with open(path, encoding="utf-8") as f:
            results.append(json.load(f))
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá trùng tên cột ở đầu file.

    eval_scores: nội dung eval_result.json (khoá "accuracy", "macro_f1"); chỉ truyền cho baseline và
    cấu hình cuối cùng. Thông tin không có cột riêng (scheduler, epoch phân kỳ, bước FP16 bị bỏ) được
    nối vào notes.
    """
    cfg, s = result["cfg"], result["summary"]
    auto_notes = []
    if cfg.get("scheduler"):
        auto_notes.append(f"scheduler={cfg['scheduler']}")
    if s.get("diverged"):
        auto_notes.append(f"NaN/inf ở epoch {s.get('diverged_epoch')}; số liệu (nếu có) lấy từ các epoch trước đó")
    if s.get("amp_skipped_steps"):
        auto_notes.append(f"GradScaler bỏ qua {s['amp_skipped_steps']} bước do gradient không hữu hạn")
    if s.get("device") == "mps":
        auto_notes.append("chạy trên Apple MPS: peak_mem_MB là mức cấp phát lớn nhất đo ở cuối mỗi epoch")
    elif s.get("device") == "cpu":
        auto_notes.append("chạy trên CPU: không đo peak_mem_MB")
    h = result.get("history") or {}
    if h.get("epoch"):
        auto_notes.append(f"grad_norm trước clip: TB các epoch {sum(h['grad_norm']) / len(h['grad_norm']):.3f}, "
                          f"bước lớn nhất {max(h['grad_norm_max']):.3g}")
        if cfg["clip_norm"] is not None:
            auto_notes.append(f"tỉ lệ bước bị clip {100 * sum(h['clip_frac']) / len(h['clip_frac']):.0f}%")
    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg["group"],
        "description": cfg["description"],
        "loss": cfg["loss"].upper(),
        "optimizer": OPTIMIZER_NAMES[cfg["optimizer"]],
        "lr": cfg["lr"],
        "weight_decay": cfg["weight_decay"],
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": "-".join(str(h) for h in cfg["hidden"]),
        "dropout": cfg["dropout"],
        "clip_norm": "none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
        "diverged": "Y" if s["diverged"] else "N",
        "eval_acc": eval_scores["accuracy"] if eval_scores else None,
        "eval_macro_f1": eval_scores["macro_f1"] if eval_scores else None,
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": "; ".join([n for n in [notes] + auto_notes if n]),
    }
    for key in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
                "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        row[key] = s.get(key)
    return row


def _cell_value(v):
    """Giá trị ghi vào ô: NaN/inf và None thành ô trống (để công thức IF(...="","",...) không lỗi)."""
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return None
    return v


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    seed_ids      : exp_id của các lần chạy baseline khác seed -> cột A sheet "Seeds" (tối đa 5)
    summary_notes : {group: nhận xét ngắn} -> cột H sheet "Summary"
    Các cột công thức (FORMULA_COLUMNS) và công thức của Seeds/Summary được giữ nguyên. openpyxl không
    tính công thức, nên sau khi lưu, giá trị của từng ô công thức được tính ở đây và ghi kèm vào file
    (_cache_formula_values): file đọc được số ngay cả khi chưa mở bằng Excel; Excel vẫn tự tính lại khi mở.
    """
    import openpyxl

    assert len(rows) <= MAX_ROWS, f"mẫu chỉ có công thức cho {MAX_ROWS} dòng, hiện {len(rows)}"
    ids = [r["exp_id"] for r in rows]
    assert len(ids) == len(set(ids)), "exp_id bị trùng"

    wb = openpyxl.load_workbook(template_path)  # không data_only: giữ công thức
    ws = wb["Experiments"]
    col_of = {cell.value: cell.column for cell in ws[1] if cell.value}
    missing = [c for c in COLUMNS if c not in col_of]
    assert not missing, f"mẫu thiếu cột: {missing}"

    # xoá giá trị điền sẵn (dòng baseline mẫu) ở các cột nhập, không đụng cột công thức
    for r in range(2, 2 + MAX_ROWS):
        for name in COLUMNS:
            ws.cell(row=r, column=col_of[name]).value = None
    for i, row in enumerate(rows):
        for name in COLUMNS:
            ws.cell(row=2 + i, column=col_of[name]).value = _cell_value(row.get(name))

    if seed_ids is not None:
        assert len(seed_ids) <= 5, "sheet Seeds chỉ có 5 dòng nhập (A2:A6)"
        ws_seeds = wb["Seeds"]
        for i in range(5):
            ws_seeds.cell(row=2 + i, column=1).value = seed_ids[i] if i < len(seed_ids) else None

    if summary_notes:
        ws_sum = wb["Summary"]
        for r in range(2, ws_sum.max_row + 1):
            group = ws_sum.cell(row=r, column=1).value
            if group in summary_notes:
                ws_sum.cell(row=r, column=8).value = summary_notes[group]

    sheet_files = {name: f"xl/worksheets/sheet{i + 1}.xml" for i, name in enumerate(wb.sheetnames)}
    summary_groups = {r: ws_sum_cell.value for r, ws_sum_cell in
                      ((r, wb["Summary"].cell(row=r, column=1)) for r in range(2, 12))}
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    _cache_formula_values(out_path, rows, seed_ids if seed_ids is not None else [], sheet_files, summary_groups)


def _formula_values(rows: list[dict], seed_ids: list[str], summary_groups: dict[int, str]) -> dict[str, dict[str, object]]:
    """Tính lại bằng Python đúng các công thức của mẫu. Trả về {sheet: {toạ độ ô: giá trị}}; "" = ô rỗng."""
    num = lambda v: _cell_value(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
    by_id = {r["exp_id"]: r for r in rows}

    # Seeds: B/C/D = val_acc / val_macro_f1 / best_val_loss của từng exp_id; dòng 8-10 = mean, std mẫu, 2σ
    seeds, stat = {}, {}
    for col, key in (("B", "val_acc"), ("C", "val_macro_f1"), ("D", "best_val_loss")):
        vals = []
        for i in range(5):
            v = num(by_id[seed_ids[i]].get(key)) if i < len(seed_ids) and seed_ids[i] in by_id else None
            seeds[f"{col}{2 + i}"] = "" if v is None else v
            if v is not None:
                vals.append(v)
        mean = statistics.fmean(vals) if vals else ""
        std = statistics.stdev(vals) if len(vals) >= 2 else ""
        seeds[f"{col}8"], seeds[f"{col}9"], seeds[f"{col}10"] = mean, std, "" if std == "" else 2 * std
        stat[key] = (mean, seeds[f"{col}10"])
    base_f1, noise = stat["val_macro_f1"]

    # Experiments: AD = step0 − ln 7, AE = val − train loss cuối, AF = Δ macro-F1 so với TB baseline, AG = vượt 2σ?
    exps = {}
    for i in range(MAX_ROWS):
        r, row = 2 + i, rows[i] if i < len(rows) else {}
        step0, tr, va, f1 = (num(row.get(k)) for k in ("step0_loss", "final_train_loss", "final_val_loss", "val_macro_f1"))
        delta = "" if f1 is None or base_f1 == "" else f1 - base_f1
        exps[f"AD{r}"] = "" if step0 is None else step0 - math.log(7)
        exps[f"AE{r}"] = "" if tr is None or va is None else va - tr
        exps[f"AF{r}"] = delta
        exps[f"AG{r}"] = "" if delta == "" or noise == "" else ("Có" if abs(delta) > noise else "Không")

    # Summary: theo nhóm
    summ, tried = {}, 0
    for r, group in summary_groups.items():
        in_group = [row for row in rows if row["group"] == group]
        f1s = [num(row.get("val_macro_f1")) for row in in_group if num(row.get("val_macro_f1")) is not None]
        accs = [num(row.get("val_acc")) for row in in_group if num(row.get("val_macro_f1")) is not None and num(row.get("val_acc")) is not None]
        summ[f"B{r}"], summ[f"C{r}"] = len(in_group), len(f1s)
        summ[f"D{r}"] = max(f1s) if f1s else ""
        summ[f"E{r}"] = min(f1s) if f1s else ""
        summ[f"F{r}"] = max(accs) if accs else ""
        if group not in ("baseline", "final", "other"):   # 7 chủ đề; ba nhóm còn lại là ô chữ "—" trong mẫu
            summ[f"G{r}"] = "Có" if f1s else "Chưa"
            tried += bool(f1s)
    summ["D13"] = tried
    return {"Experiments": exps, "Seeds": seeds, "Summary": summ}


def _cache_formula_values(xlsx_path: str, rows, seed_ids, sheet_files: dict[str, str], summary_groups) -> None:
    """Ghi giá trị đã tính vào thẻ <v> của từng ô công thức trong file .xlsx (giữ nguyên thẻ <f>)."""
    values = _formula_values(rows, seed_ids, summary_groups)
    tmp = str(xlsx_path) + ".tmp"
    with zipfile.ZipFile(xlsx_path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            sheet = next((name for name, f in sheet_files.items() if f == item.filename), None)
            if sheet in values:
                xml = data.decode("utf-8")
                for coord, value in values[sheet].items():
                    pattern = re.compile(rf'<c r="{coord}"(?P<attrs>[^>]*)>(?P<f><f>.*?</f>)<v ?/?>(?:</v>)?</c>', re.S)
                    m = pattern.search(xml)
                    assert m, f"không tìm thấy ô công thức {sheet}!{coord}"
                    attrs = re.sub(r'\s+t="[^"]*"', "", m.group("attrs"))
                    if isinstance(value, str):
                        attrs += ' t="str"'
                        cell = f'<c r="{coord}"{attrs}>{m.group("f")}<v>{escape(value)}</v></c>'
                    else:
                        cell = f'<c r="{coord}"{attrs}>{m.group("f")}<v>{value!r}</v></c>'
                    xml = xml[:m.start()] + cell + xml[m.end():]
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, xlsx_path)
